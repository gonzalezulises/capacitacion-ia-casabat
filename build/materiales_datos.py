# -*- coding: utf-8 -*-
"""Genera 22_que_puedo_subir.docx: qué se sube a una IA y qué no.

Es un BORRADOR para que CasaBat lo valide con TI y Legal, no una norma. La Ley
81 de 2019 regula el tratamiento de datos personales en Panamá, y los otros tres
países donde opera la empresa tienen sus propias obligaciones. Nada de esto se
afirma como criterio cerrado: se propone para que la empresa lo confirme.

El laboratorio 8 usa la fila del comprobante.
"""
from pathlib import Path

from docx.shared import Pt, RGBColor

from materiales_bloque4 import base_document, save

ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "materiales"

CASOS = [
    {
        "caso": "Captura de una conversación de WhatsApp con un cliente",
        "necesita": "El reclamo: qué producto, cuándo lo compró y qué pide.",
        "sobra": "El número de teléfono, la foto de perfil, el nombre completo y el resto de "
                 "la conversación que no viene al caso.",
        "antes": "Recortar la captura o tapar el número y el nombre. Basta con «cliente de "
                 "Vía España».",
        "quien": "La jefatura del área decide; TI confirma en qué cuenta puede subirse.",
    },
    {
        "caso": "Foto o copia de una factura de venta",
        "necesita": "Fecha, sucursal, producto, código y montos.",
        "sobra": "Nombre del cliente, cédula o RUC, dirección y teléfono.",
        "antes": "Tapar los datos del comprador. Para resolver una garantía casi nunca hacen "
                 "falta.",
        "quien": "Jefatura del área, con el criterio que confirme Legal.",
    },
    {
        "caso": "Cartera de clientes o listado de cobranza",
        "necesita": "Importes, antigüedad y estado, para priorizar.",
        "sobra": "Nombres, identificaciones y contactos. Es el caso de mayor riesgo de los "
                 "cuatro.",
        "antes": "Sustituir el nombre por un código antes de subir nada, y conservar la "
                 "equivalencia fuera de la herramienta.",
        "quien": "No se sube sin visto bueno de TI y Legal. Aquí la decisión no es del área.",
    },
    {
        "caso": "Fotografía del interior de una sucursal",
        "necesita": "Lo que se quiere revisar: la exhibición, el orden, la señalización.",
        "sobra": "Caras de clientes o de personal, pantallas encendidas con datos, documentos "
                 "sobre el mostrador.",
        "antes": "Repetir la foto sin personas ni pantallas. Es más rápido que recortarla "
                 "después.",
        "quien": "Quien toma la foto, con la regla escrita a la vista.",
    },
]

PREGUNTAS = [
    "¿Qué necesita de verdad la tarea? Casi siempre es menos de lo que trae el documento.",
    "¿Qué se puede quitar sin que el trabajo deje de funcionar?",
    "¿En qué cuenta y en qué herramienta puede procesarse? No es lo mismo una cuenta "
    "personal que una de la empresa con acuerdo de tratamiento.",
    "¿Quién autoriza? Si no hay nadie que pueda decir que sí, la respuesta es que no.",
    "¿Queda registro de qué se subió y para qué?",
]


def build():
    doc = base_document(
        "Qué puedo subir a una IA y qué no",
        "BORRADOR para discutir con TI y Legal de Casa de las Baterías. No es la política de "
        "la empresa y no sustituye el criterio legal. Se usa en la sesión 1, laboratorio 8, "
        "como ejercicio de decisión.")

    p = doc.add_paragraph()
    aviso = p.add_run(
        "Este documento propone un criterio; no lo establece. La Ley 81 de 2019 regula el "
        "tratamiento de datos personales en Panamá, y Costa Rica, El Salvador y Guatemala "
        "tienen sus propias obligaciones. La matriz definitiva la confirman TI y Legal, por "
        "país.")
    aviso.bold = True

    doc.add_heading("Las cinco preguntas, antes de subir nada", level=1)
    for pregunta in PREGUNTAS:
        doc.add_paragraph(pregunta, style="List Bullet")

    doc.add_heading("Cuatro casos que aparecen cada semana", level=1)
    for c in CASOS:
        doc.add_heading(c["caso"], level=2)
        for rotulo, clave in (("Qué necesita la tarea", "necesita"),
                              ("Qué sobra", "sobra"),
                              ("Qué hacer antes de subirlo", "antes"),
                              ("Quién lo decide", "quien")):
            p = doc.add_paragraph()
            p.add_run(f"{rotulo}: ").bold = True
            p.add_run(c[clave])

    doc.add_heading("Lo que este documento no resuelve", level=1)
    for punto in [
        "Qué herramientas están aprobadas en CasaBat y con qué cuenta se entra a cada una.",
        "Cuánto tiempo guarda cada proveedor lo que se sube, y si lo usa para entrenar.",
        "Las diferencias entre Panamá, Costa Rica, El Salvador y Guatemala.",
        "Qué pasa con lo que ya se subió antes de tener una regla escrita.",
    ]:
        doc.add_paragraph(punto, style="List Bullet")

    p = doc.add_paragraph()
    p.add_run("Referencia: ").bold = True
    enlace = p.add_run("Ley 81 de 2019 sobre protección de datos personales · antai.gob.pa")
    enlace.font.size = Pt(9)
    enlace.font.color.rgb = RGBColor(0x1C, 0x5C, 0x92)

    save(doc, MATERIALS / "22_que_puedo_subir.docx")


if __name__ == "__main__":
    build()
    print(f"22_que_puedo_subir.docx: {len(CASOS)} casos, {len(PREGUNTAS)} preguntas")
