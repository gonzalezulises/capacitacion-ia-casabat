# -*- coding: utf-8 -*-
"""Genera 17_tecnicas_y_cuando_usarlas.docx para la sesión 3.

Es el material de referencia del bloque 1. Cada técnica lleva quién la
recomienda, cuándo se usa, cuándo no y un caso de Casa de las Baterías.

Las tres fuentes se consultaron el 1 de octubre de 2026 y están citadas con su
URL en el documento. verifica-tecnicas.mjs comprueba que las técnicas que el
deck nombra estén fichadas aquí y que las tres fuentes sigan citadas.
"""
from pathlib import Path

from docx.shared import Pt, RGBColor

from materiales_bloque4 import base_document, save

ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "materiales"

FUENTES = [
    ("Google · Gemini API",
     "Prompt design strategies",
     "https://ai.google.dev/gemini-api/docs/prompting-strategies"),
    ("OpenAI",
     "Prompt engineering",
     "https://developers.openai.com/api/docs/guides/prompt-engineering"),
    ("Anthropic",
     "Prompting best practices",
     "https://platform.claude.com/docs/en/build-with-claude/prompt-engineering/"
     "claude-prompting-best-practices"),
]

COINCIDEN = [
    ("Di exactamente qué quieres",
     "Las tres abren por aquí. Anthropic lo resume en una regla que sirve de prueba: "
     "enséñale tu pedido a un compañero que no sepa del tema. Si él se confunde, la IA "
     "también."),
    ("Pon ejemplos",
     "Google dice que los incluyas siempre. OpenAI pide que sean variados. Anthropic "
     "concreta: de tres a cinco, parecidos a tu caso real y distintos entre sí."),
    ("Separa las partes del pedido",
     "Las tres recomiendan marcar dónde empieza la instrucción, dónde el contexto y dónde "
     "los datos. Con etiquetas o con títulos; da igual, mientras sea siempre igual."),
    ("Explica por qué",
     "No solo la orden: el motivo. Anthropic lo dice claro: el modelo generaliza a partir "
     "de la explicación. «No uses puntos suspensivos» funciona peor que «esto lo va a leer "
     "un lector automático en voz alta y no sabe pronunciarlos»."),
]

TECNICAS = [
    {
        "nombre": "1. Instrucción explícita",
        "que_es": "Decir el resultado que quieres, el formato y los límites, en vez de "
                  "describir el tema y esperar.",
        "cuando": "Siempre. Es la base de las otras ocho.",
        "cuando_no": "Nunca sobra, pero no arregla un pedido al que le falta el dato. Si no "
                     "diste el número de cotización, ninguna redacción lo inventa bien.",
        "quien": "Las tres. Anthropic: trátalo como a alguien brillante que entró ayer y no "
                 "conoce tus procesos.",
        "casabat": "En vez de «analiza el archivo de ventas», pide: «dime qué tres líneas "
                   "de producto concentran más ingreso, con el porcentaje y las filas de "
                   "donde sale».",
    },
    {
        "nombre": "2. Ejemplos (few-shot)",
        "que_es": "Pegar de tres a cinco casos resueltos antes de pedir el tuyo. El modelo "
                  "copia el patrón: el formato, el tono y hasta el criterio.",
        "cuando": "Cuando el resultado tiene una forma fija: clasificar, normalizar nombres, "
                  "escribir con tu estilo, rellenar una ficha siempre igual.",
        "cuando_no": "Si tus ejemplos se parecen demasiado entre sí, aprende el parecido y "
                     "falla en todo lo demás. Anthropic insiste en que sean diversos.",
        "quien": "Google: inclúyelos siempre. OpenAI: que cubran entradas variadas. "
                 "Anthropic: de tres a cinco, y envueltos en una etiqueta para que no se "
                 "confundan con la instrucción.",
        "casabat": "El archivo de ventas trae la misma línea escrita de cinco formas. Con "
                   "tres ejemplos de equivalencia resueltos a mano, la IA normaliza las "
                   "demás sin inventar categorías nuevas.",
    },
    {
        "nombre": "3. Tu estilo, con ejemplos",
        "que_es": "La misma técnica anterior aplicada a cómo escribes. Pegas tres textos "
                  "tuyos y pides la ficha de tu estilo antes de redactar nada.",
        "cuando": "Cuando lo que salga va a salir con tu firma y tiene que sonar a ti.",
        "cuando_no": "No sirve pedir «que no suene a IA». Eso es una queja, no una "
                     "instrucción: no dice a qué tiene que sonar.",
        "quien": "Es few-shot. Anthropic añade que con una sola frase de rol ya cambia el "
                 "tono, y OpenAI reserva un bloque llamado «Identity» para eso.",
        "casabat": "Tres textos tuyos: un correo al equipo, un informe corto y un mensaje a "
                   "un proveedor. De ahí sale cómo arrancas, qué largo tienen tus frases y "
                   "cómo das una mala noticia.",
    },
    {
        "nombre": "4. Delimitadores",
        "que_es": "Marcar con etiquetas o títulos qué parte es la instrucción, qué parte el "
                  "contexto y qué parte los datos pegados.",
        "cuando": "En cuanto el pedido mezcla instrucción con un texto largo pegado. Sin "
                  "separación, la IA confunde tus datos con órdenes.",
        "cuando_no": "En un pedido de dos líneas sobra.",
        "quien": "Las tres. OpenAI propone cuatro secciones fijas: identidad, instrucciones, "
                 "ejemplos y contexto. Anthropic y Google usan etiquetas tipo <datos>.",
        "casabat": "Al pegar la política de garantía y el correo de un cliente en el mismo "
                   "pedido, cada uno va en su etiqueta. Si no, el reclamo del cliente puede "
                   "leerse como una instrucción.",
    },
    {
        "nombre": "5. Orden en documentos largos",
        "que_es": "Poner los documentos arriba del todo y la pregunta al final, no al revés.",
        "cuando": "Con documentos largos o varios a la vez: un procedimiento con sus anexos, "
                  "un expediente completo.",
        "cuando_no": "Con un texto corto da igual el orden.",
        "quien": "Anthropic. Lo tienen medido: la pregunta al final mejora la respuesta hasta "
                 "un treinta por ciento en pruebas con varios documentos.",
        "casabat": "Procedimiento vigente, borrador y anexo de aprobación arriba, cada uno "
                   "con su nombre de archivo. La pregunta sobre qué umbral manda hoy, al "
                   "final.",
    },
    {
        "nombre": "6. Que cite antes de responder",
        "que_es": "Pedirle que primero copie los fragmentos en los que se va a apoyar, y "
                  "solo después responda.",
        "cuando": "Siempre que la respuesta tenga que sostenerse en un documento: políticas, "
                  "procedimientos, contratos, actas.",
        "cuando_no": "En tareas de redacción libre, donde no hay fuente que citar.",
        "quien": "Anthropic lo recomienda para documentos largos: ayuda a centrarse en lo "
                 "que importa e ignorar el resto.",
        "casabat": "Antes de decirte si una batería de moto de ocho meses tiene reemplazo, "
                   "que copie el párrafo de la política que lo dice.",
    },
    {
        "nombre": "7. Pide las columnas exactas",
        "que_es": "Decir las columnas que quieres y en qué orden, en vez de aceptar un "
                  "texto corrido que luego hay que desarmar a mano.",
        "cuando": "Cuando el resultado va a alimentar otra cosa: una hoja de cálculo, un "
                  "tablero, una revisión fila por fila.",
        "cuando_no": "Cuando lo que necesitas es un texto para leer. Una tabla no es un "
                     "correo.",
        "quien": "Las tres. Anthropic lo resume en una regla útil: dile el formato que "
                 "quieres, no el que no quieres. En el chat esto se pide por escrito; los "
                 "esquemas que menciona la documentación son del API.",
        "casabat": "La auditoría de los 41 procedimientos sale en columnas fijas: archivo, "
                   "regla que rompe, parte mala y nombre corregido. Así se pega en la hoja "
                   "y se revisa.",
    },
    {
        "nombre": "8. Partir y encadenar",
        "que_es": "Dividir un trabajo largo en pasos, y que cada paso tome el resultado del "
                  "anterior.",
        "cuando": "Cuando necesitas ver el resultado intermedio y aprobarlo antes de seguir. "
                  "La cadena más útil es: borrador, revisión contra criterios, versión final.",
        "cuando_no": "Los modelos de hoy resuelven solos buena parte de los pasos. Encadena "
                     "cuando quieras inspeccionar el intermedio, no por costumbre.",
        "quien": "Google: un pedido por instrucción, o encadenados en secuencia. Anthropic "
                 "nombra la autocorrección como el patrón más común. OpenAI pide planificar "
                 "y dejar ver el avance.",
        "casabat": "Paso uno: extraer los datos del expediente. Paso dos: revisarlos contra "
                   "la regla de nombres. Paso tres: redactar la lista de correcciones. Se "
                   "revisa el paso dos antes de redactar nada.",
    },
    {
        "nombre": "9. Dale un rol",
        "que_es": "Decirle desde qué puesto trabaja: control documental, jefatura "
                  "administrativa, analista comercial.",
        "cuando": "Cuando el mismo dato se mira distinto según quién lo mire.",
        "cuando_no": "No hace falta inventarle una biografía. Una frase basta.",
        "quien": "Anthropic: una sola frase de rol ya cambia el resultado. OpenAI lo pone en "
                 "su bloque de identidad.",
        "casabat": "«Eres quien revisa el control documental de Casa de las Baterías» da una "
                   "revisión distinta a «eres el dueño del proceso».",
    },
]

NO_HACER = [
    ("Pedir instrucciones cada vez más detalladas",
     "OpenAI distingue: a un modelo de razonamiento se le da la idea general, como a un "
     "compañero con experiencia. Al otro tipo, instrucciones precisas, como a alguien que "
     "entró ayer. Más detalle no siempre es mejor."),
    ("Pedirle que te enseñe cómo piensa",
     "Pide el método, la fuente y las comprobaciones. Eso sí se puede revisar."),
    ("Tratar lo que viene en un archivo como una orden",
     "Un correo o un documento pegado son datos, no instrucciones. Si el texto trae algo "
     "parecido a una orden, que lo reporte y no lo obedezca."),
    ("Copiar ajustes del API al chat",
     "La temperatura, los esquemas de salida y el presupuesto de razonamiento viven en el "
     "API. En la ventana de chat no están: ahí se consigue lo mismo con instrucción, "
     "ejemplos y formato pedido por escrito."),
]

DECISION = [
    ("Tengo que clasificar u ordenar muchas filas iguales", "Ejemplos · columnas exactas"),
    ("Tiene que sonar a mí", "Tu estilo con ejemplos · rol"),
    ("La respuesta depende de un documento", "Delimitadores · orden · que cite antes"),
    ("Son varios documentos largos", "Datos arriba, pregunta al final · que cite antes"),
    ("El resultado se pega en una hoja", "Pide las columnas exactas"),
    ("Es un trabajo de varios pasos y quiero revisar en medio", "Partir y encadenar"),
    ("Salió casi bien y solo quiero un ajuste", "Instrucción explícita de qué no se toca"),
    ("Hay que calcular", "Pide la fórmula y las filas; comprueba dos cifras a mano"),
]


def build():
    doc = base_document(
        "Técnicas de prompting y cuándo usar cada una",
        "Material de consulta de la sesión 3. Recoge lo que documentan Google, OpenAI y "
        "Anthropic, consultado el 1 de octubre de 2026, con un caso de Casa de las Baterías "
        "por técnica. Los casos son ficticios.")

    doc.add_heading("De dónde sale esto", level=1)
    doc.add_paragraph(
        "Tres fuentes oficiales, no artículos de terceros. Si una recomendación de este "
        "documento te choca, ve a la fuente: las interfaces cambian cada pocos meses.")
    for casa, titulo, url in FUENTES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{casa} · ").bold = True
        p.add_run(f"{titulo}. ")
        enlace = p.add_run(url)
        enlace.font.size = Pt(9)
        enlace.font.color.rgb = RGBColor(0x1C, 0x5C, 0x92)

    doc.add_heading("En qué coinciden las tres", level=1)
    doc.add_paragraph(
        "Antes de las diferencias, lo que las tres dicen igual. Si solo te llevas esto, ya "
        "cambia el resultado.")
    for titulo, texto in COINCIDEN:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{titulo}. ").bold = True
        p.add_run(texto)

    doc.add_heading("Las nueve técnicas", level=1)
    doc.add_paragraph(
        "Cada una con lo mismo: qué es, cuándo se usa, cuándo no, quién la recomienda y un "
        "caso de la empresa. La columna de «cuándo no» es la que más se salta y la que más "
        "tiempo ahorra.")
    for t in TECNICAS:
        doc.add_heading(t["nombre"], level=2)
        for rotulo, clave in (("Qué es", "que_es"), ("Cuándo", "cuando"),
                              ("Cuándo no", "cuando_no"), ("Quién lo dice", "quien"),
                              ("En CasaBat", "casabat")):
            p = doc.add_paragraph()
            p.add_run(f"{rotulo}: ").bold = True
            p.add_run(t[clave])

    doc.add_heading("Lo que ya no se recomienda", level=1)
    for titulo, texto in NO_HACER:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{titulo}. ").bold = True
        p.add_run(texto)

    doc.add_heading("Si tu tarea es… usa", level=1)
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Table Grid"
    encabezado = tabla.rows[0].cells
    encabezado[0].text = "Tu tarea"
    encabezado[1].text = "Técnicas"
    for celda in encabezado:
        for run in celda.paragraphs[0].runs:
            run.font.bold = True
    for tarea, tecnica in DECISION:
        fila = tabla.add_row().cells
        fila[0].text = tarea
        fila[1].text = tecnica

    doc.add_heading("Una advertencia sobre las fuentes", level=1)
    doc.add_paragraph(
        "Buena parte de esta documentación está escrita para quien programa contra el API, "
        "no para quien usa la ventana de chat. Lo que aquí aparece se eligió porque funciona "
        "igual en el chat. Cuando algo solo existe en el API, está dicho.")

    save(doc, MATERIALS / "17_tecnicas_y_cuando_usarlas.docx")


if __name__ == "__main__":
    build()
    print(f"17_tecnicas_y_cuando_usarlas.docx: {len(TECNICAS)} técnicas, "
          f"{len(FUENTES)} fuentes, {len(DECISION)} filas de decisión")
