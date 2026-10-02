# -*- coding: utf-8 -*-
"""Genera 18_reclamos_de_clientes.docx.

Cinco reclamos ficticios, escritos como los escribiría un cliente: con lo que le
importa delante y el dato que decide escondido a mitad de párrafo. Ninguno dice
si procede o no; eso es el ejercicio.

El laboratorio 5 de la sesión 3 prometía «el reclamo del cliente» y ese texto no
existía en ninguna parte. verifica-materiales.mjs comprueba ahora que todo
insumo prometido por un laboratorio exista de verdad.
"""
from pathlib import Path

from docx.shared import Pt, RGBColor

from materiales_bloque4 import base_document, save

ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "materiales"

RECLAMOS = [
    {
        "id": "R-01",
        "via": "Correo a la sucursal de Vía España",
        "asunto": "Batería de moto fallada — compra de julio",
        "texto": "Buenos días. Les compramos una batería de moto el 14 de julio y ya no "
                 "levanta carga. Tiene tres meses de uso, así que está dentro de la garantía "
                 "de seis. La tenemos puesta en el montacargas chico del depósito, que es "
                 "para lo que la necesitábamos, y hasta la semana pasada trabajó bien. "
                 "Adjunto la factura. Necesito el reemplazo esta semana porque el equipo "
                 "está parado.",
    },
    {
        "id": "R-02",
        "via": "WhatsApp a la sucursal de Tocumen",
        "asunto": "Batería de auto que no arranca",
        "texto": "Buenas, compré una batería para mi carro el año pasado en noviembre y ya "
                 "está fallando. Me dijeron que tenía garantía. La batería es de la línea "
                 "normal, no la cara. ¿Me la cambian?",
    },
    {
        "id": "R-03",
        "via": "Correo del encargado de flota",
        "asunto": "Seis baterías de la flota — garantía",
        "texto": "Estimados: de las doce baterías que nos instalaron en marzo, seis están "
                 "presentando fallas. No tengo a mano las facturas porque el compañero que "
                 "llevaba eso ya no está en la empresa, pero ustedes deben tener el registro "
                 "de la compra. Les paso las placas de los vehículos. Quedo atento al cambio "
                 "de las seis.",
    },
    {
        "id": "R-04",
        "via": "Llamada registrada en la sucursal de David",
        "asunto": "Quedó mal instalada",
        "texto": "El cliente llama molesto. Dice que la batería se la instalaron a domicilio "
                 "hace mes y medio y que desde entonces el carro hace un ruido raro al "
                 "arrancar. No reclama la batería: dice que el problema es cómo quedaron los "
                 "bornes. Pide que vayan a revisarla sin costo.",
    },
    {
        "id": "R-05",
        "via": "Correo a servicio al cliente",
        "asunto": "Reclamo batería con dos meses",
        "texto": "Les escribo porque la batería que compré hace dos meses está derramando. "
                 "Aclaro que el carro se me volcó en un accidente el mes pasado y la batería "
                 "quedó golpeada, pero el taller me dijo que eso no debería hacer que "
                 "derrame. Me parece que les vendieron un producto defectuoso. Espero su "
                 "respuesta.",
    },
]


def build():
    doc = base_document(
        "Cinco reclamos de clientes",
        "Material de práctica del taller. Reclamos ficticios, escritos como llegan de "
        "verdad: con el pedido por delante y el dato que decide a mitad de párrafo. "
        "Ninguno indica si procede. Se usan en la sesión 3, laboratorio 5.")

    doc.add_paragraph(
        "Cada reclamo se contrasta contra 03_politica_garantia.docx. El dato que resuelve el "
        "caso casi nunca es el que el cliente pone en primer plano.")

    for r in RECLAMOS:
        doc.add_heading(f"{r['id']} · {r['asunto']}", level=1)
        p = doc.add_paragraph()
        via = p.add_run(r["via"])
        via.italic = True
        via.font.size = Pt(9.5)
        via.font.color.rgb = RGBColor(70, 70, 70)
        doc.add_paragraph(r["texto"])

    save(doc, MATERIALS / "18_reclamos_de_clientes.docx")


if __name__ == "__main__":
    build()
    print(f"18_reclamos_de_clientes.docx: {len(RECLAMOS)} reclamos")
