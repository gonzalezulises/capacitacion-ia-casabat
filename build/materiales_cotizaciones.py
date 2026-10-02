# -*- coding: utf-8 -*-
"""Genera 21_cotizaciones_semestre.xlsx: el embudo comercial del semestre.

Sesenta cotizaciones ficticias. El patrón que hay que descubrir es incómodo:
el descuento alto no mejora la conversión. Quien ordene por importe se lleva a
las equivocadas; quien cruce días sin respuesta, margen y motivo de pérdida,
encuentra otras.

Los defectos están puestos a propósito, como en el resto de los materiales:
fechas en dos formatos, un importe guardado como texto, dos filas duplicadas y
un margen imposible. verifica-pipeline.mjs los recuenta desde el archivo.
"""
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "materiales"

BLUE = "1C5C92"

CABECERAS = ["n", "cotizacion", "fecha_envio", "cliente", "pais", "canal", "linea_producto",
             "importe_usd", "descuento_pct", "margen_pct", "dias_sin_respuesta", "estado",
             "motivo_perdida", "vendedor"]

CLIENTES = [
    ("Transportes Istmo", "Panamá"), ("Flota Delta", "Panamá"),
    ("Distribuidora Chiriquí", "Panamá"), ("Taller Vía Brasil", "Panamá"),
    ("Logística Central", "Costa Rica"), ("Buses del Valle", "Costa Rica"),
    ("Agroexport Alajuela", "Costa Rica"), ("Reparto Soyapango", "El Salvador"),
    ("Industrias Santa Ana", "El Salvador"), ("Transporte Quetzal", "Guatemala"),
    ("Mixco Cargo", "Guatemala"), ("Minas del Norte", "Guatemala"),
]
CANALES = ["Sucursal", "WhatsApp", "Call center", "Visita a flota"]
LINEAS = ["Batería Auto", "Batería Moto", "Batería Industrial", "Energía Solar",
          "Respaldo de Energía"]
VENDEDORES = ["M. Ríos", "A. Pinto", "L. Serrano", "C. Vega"]
MOTIVOS = ["Precio", "Plazo de entrega", "Se fue con la competencia",
           "No respondió", "Presupuesto congelado", ""]


def build():
    random.seed(21)
    filas = []
    n = 0

    # El patrón: por encima del 12 % de descuento la conversión NO mejora.
    # Se construye a propósito para que el dato lo sostenga.
    for tramo, descuentos, prob_ganada in (("bajo", (0, 5, 8), 0.55),
                                           ("medio", (10, 12), 0.60),
                                           ("alto", (15, 18, 22), 0.28)):
        for _ in range(20):
            n += 1
            cliente, pais = random.choice(CLIENTES)
            desc = random.choice(descuentos)
            importe = round(random.uniform(900, 18500), 2)
            margen = round(max(2.0, 34 - desc * 1.4 + random.uniform(-3, 3)), 1)
            ganada = random.random() < prob_ganada
            dias = random.choice([2, 4, 7, 9, 12, 15, 18, 23, 31, 46])
            if ganada:
                estado, motivo, dias = "Ganada", "", random.choice([1, 2, 3, 5])
            elif dias >= 15:
                estado = "Vencida"
                motivo = random.choice(["No respondió", "Presupuesto congelado"])
            else:
                estado = random.choice(["Abierta", "Perdida"])
                motivo = "" if estado == "Abierta" else random.choice(MOTIVOS[:3])
            mes = random.randint(1, 6)
            dia = random.randint(1, 28)
            # dos formatos de fecha, como en el archivo de ventas
            fecha = (f"2026-{mes:02d}-{dia:02d}" if n % 7 else f"{dia:02d}/{mes:02d}/2026")
            filas.append([n, f"COT-2026-{3000 + n}", fecha, cliente, pais,
                          random.choice(CANALES), random.choice(LINEAS), importe, desc,
                          margen, dias, estado, motivo, random.choice(VENDEDORES)])

    # --- defectos puestos a propósito ---
    filas[11][7] = f"{filas[11][7]:,.2f}"          # importe guardado como texto
    filas[29][9] = 118.0                            # margen imposible
    filas[44][10] = ""                              # días sin respuesta en blanco
    filas.append(list(filas[7]));  filas[-1][0] = len(filas)   # duplicado
    filas.append(list(filas[23])); filas[-1][0] = len(filas)   # duplicado

    wb = Workbook()
    ws = wb.active
    ws.title = "Cotizaciones"
    for col, nombre in enumerate(CABECERAS, start=1):
        c = ws.cell(row=1, column=col, value=nombre)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    borde = Side(style="thin", color="D3DEEA")
    for i, fila in enumerate(filas, start=2):
        for col, valor in enumerate(fila, start=1):
            c = ws.cell(row=i, column=col, value=valor)
            c.font = Font(size=9.5)
            c.border = Border(bottom=borde)
            if i % 2:
                c.fill = PatternFill("solid", fgColor="F7FAFD")

    for col, ancho in enumerate([4, 16, 13, 24, 13, 14, 19, 13, 14, 12, 13, 11, 25, 12], start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = ws.cell(row=2, column=1)

    leeme = wb.create_sheet("Léeme")
    leeme["A1"] = "Casa de las Baterías · Cotizaciones del semestre"
    leeme["A1"].font = Font(bold=True, size=14, color=BLUE)
    for i, aviso in enumerate([
        "Datos ficticios. Corte del 30 de junio de 2026.",
        "Ningún cliente, cotización, vendedor ni importe de este archivo es real.",
        "El archivo trae defectos puestos a propósito: fechas en dos formatos, un importe "
        "guardado como texto, un margen imposible, una celda vacía y dos filas duplicadas.",
        "Se usa en la sesión 1, laboratorio 7.",
    ], start=3):
        leeme.cell(row=i, column=1, value=aviso).alignment = Alignment(wrap_text=True)
    leeme.column_dimensions["A"].width = 96

    MATERIALS.mkdir(parents=True, exist_ok=True)
    wb.save(MATERIALS / "21_cotizaciones_semestre.xlsx")
    return len(filas)


if __name__ == "__main__":
    total = build()
    print(f"21_cotizaciones_semestre.xlsx: {total} filas")
